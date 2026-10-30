#!/usr/bin/env python3
"""
Computes the evidence base for the analytics question bank and the data
commentary pack: actual findings from VPT_Consumption_Dataset.xlsx.

Writes data/findings.json (all computed metrics, used by the builders).
"""
import os
import json
import calendar
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "..", "data", "VPT_Consumption_Dataset.xlsx")
OUT = os.path.join(BASE, "..", "data", "findings.json")

wb = load_workbook(SRC, read_only=True, data_only=True)


def rows(name):
    ws = wb[name]
    it = ws.iter_rows(values_only=True)
    next(it)
    return [list(r) for r in it if r[0] is not None]


projects = rows("03_Dim_Project")
items = rows("04_Dim_Item")
suppliers = rows("05_Dim_Supplier")
warehouses = rows("06_Dim_Warehouse")
pos = rows("08_Fact_PurchaseOrder")
grns = rows("09_Fact_Receipts")
issues = rows("10_Fact_Consumption")
budgets = rows("11_Fact_Budget")
service = rows("12_Fact_ServiceDelivery")

prj = {p[0]: p for p in projects}
itm = {i[0]: i for i in items}
sup = {s[0]: s for s in suppliers}

TRANSFER = "Inter-store Transfer"


def d(x):
    return x.date() if isinstance(x, datetime) else x


for r in issues:
    r[1] = d(r[1]); r[14] = d(r[14])
for r in pos:
    r[1] = d(r[1]); r[3] = d(r[3])
for r in grns:
    r[1] = d(r[1]); r[10] = d(r[10]); r[11] = d(r[11])
for r in service:
    r[1] = d(r[1])
for p in projects:
    p[7] = d(p[7]); p[8] = d(p[8])

ASAT = max(r[1] for r in issues)
START = min(r[1] for r in issues)


def month_key(dt):
    return f"{dt.year}-{dt.month:02d}"


def fy_of(dt):
    return f"FY{dt.year if dt.month >= 7 else dt.year - 1}-{str((dt.year + 1) if dt.month >= 7 else dt.year)[2:]}"


F = {}

# ---------------------------------------------------------------- core totals
billable = [r for r in issues if r[15] != TRANSFER]
cons_total = sum(r[10] for r in billable)
cons_all = sum(r[10] for r in issues)
transfers = cons_all - cons_total
po_live = [r for r in pos if r[14] != "Cancelled"]
po_total = sum(r[12] or 0 for r in po_live)
PRODUCT_CATS = ["Medicines", "Medical Consumables", "Laboratory", "Nutrition", "Equipment", "General Items"]
budget_product = sum(b[3] for b in budgets if b[2] in PRODUCT_CATS)
unattr = [r for r in billable if r[2] == "GEN-0000"]
unattr_val = sum(r[10] for r in unattr)
benef_total = sum(r[4] for r in service)

F.update({
    "as_at": ASAT.isoformat(), "start": START.isoformat(),
    "consumption": round(cons_total, 2), "consumption_all": round(cons_all, 2),
    "transfers": round(transfers, 2), "purchase": round(po_total, 2),
    "ratio": round(cons_total / po_total, 3), "budget_product": round(budget_product, 2),
    "utilisation": round(cons_total / budget_product, 3),
    "unattributed": round(unattr_val, 2), "unattributed_pct": round(unattr_val / cons_total * 100, 2),
    "unattributed_lines": len(unattr),
    "beneficiaries": benef_total,
    "cost_per_beneficiary": round(cons_total / benef_total, 2),
})

# ---------------------------------------------------------------- monthly
monthly = defaultdict(float)
monthly_qty = defaultdict(float)
monthly_po = defaultdict(float)
for r in billable:
    monthly[month_key(r[1])] += r[10]
    monthly_qty[month_key(r[1])] += r[7]
for r in po_live:
    monthly_po[month_key(r[1])] += r[12] or 0
months = sorted(monthly)
peak = max(months, key=lambda k: monthly[k])
trough = min(months, key=lambda k: monthly[k])
F["months"] = len(months)
F["peak_month"] = {"key": peak, "value": round(monthly[peak], 2),
                   "label": f"{calendar.month_abbr[int(peak[5:])]}-{peak[:4]}"}
F["trough_month"] = {"key": trough, "value": round(monthly[trough], 2),
                     "label": f"{calendar.month_abbr[int(trough[5:])]}-{trough[:4]}"}
avg_month = sum(monthly.values()) / len(monthly)
F["avg_month"] = round(avg_month, 2)
F["peak_vs_avg"] = round(monthly[peak] / avg_month - 1, 3)

latest_key = months[-1]
prior3 = months[-4:-1]
F["latest"] = {
    "label": f"{calendar.month_abbr[int(latest_key[5:])]}-{latest_key[:4]}",
    "value": round(monthly[latest_key], 2),
    "prior_3mo_avg": round(sum(monthly[k] for k in prior3) / len(prior3), 2),
    "vs_3mo_avg": round(monthly[latest_key] / (sum(monthly[k] for k in prior3) / len(prior3)) - 1, 3),
    "mom": round(monthly[latest_key] / monthly[months[-2]] - 1, 3) if monthly[months[-2]] else None,
}
F["monthly_series"] = {f"{calendar.month_abbr[int(k[5:])]}-{k[:4]}": round(monthly[k], 2) for k in months}

# seasonality index by calendar month
seas = defaultdict(list)
for k in months:
    seas[int(k[5:])].append(monthly[k])
season_index = {m: round((sum(v) / len(v)) / (sum(monthly.values()) / len(months)), 3) for m, v in seas.items()}
F["season_index"] = {calendar.month_abbr[m]: v for m, v in sorted(season_index.items())}
F["peak_season_month"] = calendar.month_abbr[max(season_index, key=season_index.get)]
F["peak_season_index"] = max(season_index.values())
F["low_season_month"] = calendar.month_abbr[min(season_index, key=season_index.get)]
F["low_season_index"] = min(season_index.values())

# fiscal year totals
fy_cons = defaultdict(float)
for r in billable:
    fy_cons[fy_of(r[1])] += r[10]
F["fy_consumption"] = {k: round(v, 2) for k, v in sorted(fy_cons.items())}
fy_prev, fy_cur = "FY2024-25", "FY2025-26"
F["yoy_growth"] = round(fy_cons[fy_cur] / fy_cons[fy_prev] - 1, 3)

# ---------------------------------------------------------------- project level
proj_val = defaultdict(float)
proj_po = defaultdict(float)
proj_lines = defaultdict(int)
for r in billable:
    proj_val[r[2]] += r[10]
    proj_lines[r[2]] += 1
for r in po_live:
    proj_po[r[4]] += r[12] or 0
named = {k: v for k, v in proj_val.items() if k != "GEN-0000"}
top_pid = max(named, key=named.get)
bottom_pid = min(named, key=named.get)
F["top_project"] = {"id": top_pid, "name": prj[top_pid][1], "value": round(named[top_pid], 2),
                    "share": round(named[top_pid] / cons_total, 3),
                    "ratio": round(named[top_pid] / proj_po[top_pid], 3) if proj_po[top_pid] else None}
F["bottom_project"] = {"id": bottom_pid, "name": prj[bottom_pid][1], "value": round(named[bottom_pid], 2),
                       "share": round(named[bottom_pid] / cons_total, 3)}
F["project_count"] = len(named)
ratios = {p: proj_val[p] / proj_po[p] for p in named if proj_po[p]}
lo_p = min(ratios, key=ratios.get)
hi_p = max(ratios, key=ratios.get)
F["lowest_ratio_project"] = {"id": lo_p, "name": prj[lo_p][1], "ratio": round(ratios[lo_p], 3)}
F["highest_ratio_project"] = {"id": hi_p, "name": prj[hi_p][1], "ratio": round(ratios[hi_p], 3)}

# per project cost per beneficiary
benef_proj = defaultdict(int)
for r in service:
    benef_proj[r[0]] += r[4]
cpb = {p: named[p] / benef_proj[p] for p in named if benef_proj.get(p)}
cpb_hi = max(cpb, key=cpb.get)
cpb_lo = min(cpb, key=cpb.get)
F["cpb"] = {"overall": F["cost_per_beneficiary"],
            "highest": {"id": cpb_hi, "name": prj[cpb_hi][1], "value": round(cpb[cpb_hi], 2)},
            "lowest": {"id": cpb_lo, "name": prj[cpb_lo][1], "value": round(cpb[cpb_lo], 2)},
            "spread": round(cpb[cpb_hi] / cpb[cpb_lo] - 1, 3)}

# projects ending soon with budget left
ending = []
for p in projects:
    days_left = (p[8] - ASAT).days
    if 0 < days_left <= 200:
        bud = sum(b[3] for b in budgets if b[0] == p[0])
        ending.append({"id": p[0], "name": p[1], "end": p[8].isoformat(), "days_left": days_left,
                       "consumption": round(proj_val.get(p[0], 0), 2), "budget": round(bud, 2),
                       "unspent": round(bud - proj_val.get(p[0], 0), 2)})
ending.sort(key=lambda x: x["days_left"])
F["projects_ending"] = ending

# ---------------------------------------------------------------- product level
item_val = defaultdict(float)
item_qty = defaultdict(float)
item_po = defaultdict(float)
for r in billable:
    item_val[r[4]] += r[10]
    item_qty[r[4]] += r[7]
for r in po_live:
    item_po[r[6]] += r[12] or 0
top_i = max(item_val, key=item_val.get)
F["top_item"] = {"id": top_i, "name": itm[top_i][1], "value": round(item_val[top_i], 2),
                 "share": round(item_val[top_i] / cons_total, 3)}
sorted_items = sorted(item_val.items(), key=lambda x: -x[1])
top10_share = sum(v for _, v in sorted_items[:10]) / cons_total
n_for_80 = next(i + 1 for i, (_, v) in enumerate(sorted_items)
                if sum(x for _, x in sorted_items[:i + 1]) / cons_total >= 0.8)
F["concentration"] = {"top10_share": round(top10_share, 3), "items_for_80": n_for_80,
                      "item_count": len(sorted_items)}

cat_val = defaultdict(float)
for r in billable:
    cat_val[r[6]] += r[10]
F["category"] = {k: {"value": round(v, 2), "share": round(v / cons_total, 3)}
                 for k, v in sorted(cat_val.items(), key=lambda x: -x[1])}

# ---------------------------------------------------------------- price movement
price = defaultdict(list)
for r in po_live:
    if r[11]:
        price[r[6]].append((r[1], r[11], r[9]))
q1_end = date(START.year, START.month, 1) + timedelta(days=180)
q4_start = ASAT - timedelta(days=180)
risers = []
for iid, h in price.items():
    early = [x for x in h if x[0] < q1_end]
    late = [x for x in h if x[0] >= q4_start]
    if not early or not late:
        continue
    e = sum(x[1] * x[2] for x in early) / sum(x[2] for x in early)
    l = sum(x[1] * x[2] for x in late) / sum(x[2] for x in late)
    if e <= 0:
        continue
    risers.append({"id": iid, "name": itm[iid][1], "early": round(e, 2), "late": round(l, 2),
                   "change": round(l / e - 1, 3)})
risers.sort(key=lambda x: -x["change"])
F["price_risers"] = risers[:5]
F["price_fallers"] = sorted(risers, key=lambda x: x["change"])[:3]
F["avg_price_change"] = round(sum(r["change"] for r in risers) / len(risers), 3) if risers else None
F["price_items_measured"] = len(risers)

# ---------------------------------------------------------------- stock and expiry
stock = {}
for r in grns:
    k = (r[4], r[9])
    if k in stock:
        stock[k]["qty"] += r[6]
    else:
        stock[k] = {"item": r[4], "batch": r[9], "wh": r[8], "qty": float(r[6]),
                    "exp": r[11], "cost": r[12] or 0.0}
for r in billable + [x for x in issues if x[15] == TRANSFER]:
    k = (r[4], r[13])
    if k in stock:
        stock[k]["qty"] -= r[7]

onhand_val = 0.0
expired_val = 0.0
expired_qty = 0.0
buckets = {"0-30": 0.0, "31-60": 0.0, "61-90": 0.0, "91-180": 0.0, ">180": 0.0}
expired_items = defaultdict(float)
onhand_item = defaultdict(float)
for k, b in stock.items():
    if b["qty"] <= 0:
        continue
    val = b["qty"] * b["cost"]
    onhand_val += val
    onhand_item[b["item"]] += b["qty"]
    if isinstance(b["exp"], date):
        days = (b["exp"] - ASAT).days
        if days < 0:
            expired_val += val
            expired_qty += b["qty"]
            expired_items[b["item"]] += val
        elif days <= 30:
            buckets["0-30"] += val
        elif days <= 60:
            buckets["31-60"] += val
        elif days <= 90:
            buckets["61-90"] += val
        elif days <= 180:
            buckets["91-180"] += val
        else:
            buckets[">180"] += val
F["stock"] = {"total_value": round(onhand_val, 2), "expired_value": round(expired_val, 2),
              "expired_qty": round(expired_qty), "expired_share": round(expired_val / onhand_val, 4),
              "near_expiry_60": round(buckets["0-30"] + buckets["31-60"], 2),
              "near_expiry_90": round(buckets["61-90"], 2),
              "buckets": {k: round(v, 2) for k, v in buckets.items()}}
if expired_items:
    ei = max(expired_items, key=expired_items.get)
    F["stock"]["worst_expired_item"] = {"name": itm[ei][1], "value": round(expired_items[ei], 2)}

# recent consumption for cover
look_start = ASAT - timedelta(days=91)
recent_item = defaultdict(float)
for r in billable:
    if r[1] >= look_start:
        recent_item[r[4]] += r[7]
cover = {}
for iid, q in onhand_item.items():
    if recent_item.get(iid):
        cover[iid] = q / (recent_item[iid] / 91.0)
critical = {k: v for k, v in cover.items() if v < 15}
low = {k: v for k, v in cover.items() if 15 <= v < 30}
slow = [i for i in onhand_item if i not in cover and onhand_item[i] > 0]
slow_val = sum(onhand_item[i] * itm[i][6] for i in slow)
over_180 = {k: v for k, v in cover.items() if v > 180}
over_180_val = sum(onhand_item[k] * itm[k][6] for k in over_180)
under30 = {k: v for k, v in cover.items() if v < 30}
under60 = {k: v for k, v in cover.items() if 30 <= v < 60}
F["cover"] = {"critical_count": len(critical), "low_count": len(low),
              "under30": len(under30), "under60": len(under60),
              "slow_items": len(slow), "slow_value": round(slow_val, 2),
              "over_180_count": len(over_180), "over_180_value": round(over_180_val, 2),
              "items_with_cover": len(cover)}
if cover:
    worst = min(cover, key=cover.get)
    F["cover"]["lowest"] = {"name": itm[worst][1], "days": round(cover[worst])}
    best = max(cover, key=cover.get)
    F["cover"]["highest"] = {"name": itm[best][1], "days": round(cover[best])}

# ---------------------------------------------------------------- commitment
recv_by_po = defaultdict(float)
for r in grns:
    if r[2] and r[2] != "OPENING":
        recv_by_po[r[2]] += r[6]
open_commit = 0.0
open_qty = 0.0
for r in po_live:
    got = recv_by_po.get(r[0], 0.0)
    rem = (r[9] or 0) - got
    if rem > 0:
        open_commit += rem * (r[11] or 0)
        open_qty += 1
F["open_commitment"] = round(open_commit, 2)
F["open_commitment_lines"] = round(open_qty)
F["cancelled_lines"] = sum(1 for r in pos if r[14] == "Cancelled")
F["cancelled_value"] = round(sum(r[12] or 0 for r in pos if r[14] == "Cancelled"), 2)

# ---------------------------------------------------------------- suppliers
sup_lead = defaultdict(list)
sup_val = defaultdict(float)
sup_pos = defaultdict(int)
for r in po_live:
    sup_pos[r[5]] += 1
    sup_val[r[5]] += r[12] or 0
    if r[0] in recv_by_po:
        sup_lead[r[5]].append((r[1], r[0]))
lead_days = {}
for sid, lst in sup_lead.items():
    diffs = []
    for pdate, po in lst:
        gdates = [g[1] for g in grns if g[2] == po]
        if gdates:
            diffs.append((min(gdates) - pdate).days)
    if diffs:
        lead_days[sid] = sum(diffs) / len(diffs)
if lead_days:
    worst_sup = max(lead_days, key=lead_days.get)
    F["supplier"] = {"avg_lead": round(sum(lead_days.values()) / len(lead_days)),
                     "worst": {"name": sup[worst_sup][1], "days": round(lead_days[worst_sup])},
                     "count": len(sup_pos),
                     "top": {"name": sup[max(sup_val, key=sup_val.get)][1],
                             "value": round(max(sup_val.values()), 2)}}
F["blacklisted_in_history"] = sum(1 for s in suppliers if s[6] == "Blacklisted")

# ---------------------------------------------------------------- budget
bud_cat = defaultdict(float)
cons_cat = defaultdict(float)
for b in budgets:
    if b[2] in PRODUCT_CATS:
        bud_cat[(b[1], b[2])] += b[3]
over, under = [], []
for r in billable:
    cons_cat[(fy_of(r[1]), r[6])] += r[10]
for k, b in bud_cat.items():
    c = cons_cat.get(k, 0)
    util = c / b if b else 0
    (over if util > 1 else under).append({"fy": k[0], "cat": k[1], "budget": round(b, 2),
                                          "consumption": round(c, 2), "util": round(util, 3)})
F["budget_over"] = sorted(over, key=lambda x: -x["util"])[:5]
F["budget_under"] = sorted(under, key=lambda x: x["util"])[:5]
F["budget_over_count"] = len(over)

# ---------------------------------------------------------------- data quality
dq_dup = defaultdict(int)
for r in grns:
    dq_dup[(r[0], r[9])] += 1
F["dq"] = {
    "duplicate_grc": sum(1 for v in dq_dup.values() if v > 1),
    "zero_price_po": sum(1 for r in pos if not r[11]),
    "unattributed_lines": len(unattr),
    "unattributed_pct": F["unattributed_pct"],
    "inactive_items": sum(1 for i in items if i[10] == "Inactive"),
    "blacklisted": F["blacklisted_in_history"],
}
ua_month = defaultdict(float)
for r in unattr:
    ua_month[month_key(r[1])] += r[10]
worst_ua = max(ua_month, key=ua_month.get)
F["dq"]["worst_month"] = {"label": f"{calendar.month_abbr[int(worst_ua[5:])]}-{worst_ua[:4]}",
                          "value": round(ua_month[worst_ua], 2),
                          "share": round(ua_month[worst_ua] / monthly[worst_ua], 3)}

# ---------------------------------------------------------------- projections
last3 = months[-3:]
run_rate = sum(monthly[k] for k in last3) / 3
F["run_rate"] = round(run_rate, 2)
F["next_q_projection"] = round(run_rate * 3, 2)
F["seasonal_next_q"] = round(sum(monthly[k] * season_index[int(k[5:])] for k in months[-12:]) / 12 * 3
                             / (sum(season_index.values()) / 12), 2)
fy27_months = [k for k in months if fy_of(date(int(k[:4]), int(k[5:]), 1)) == "FY2026-27"]
ytd = sum(monthly[k] for k in fy27_months)
months_elapsed = len(fy27_months)
F["fy2026_27"] = {"ytd": round(ytd, 2), "months": months_elapsed,
                  "burn": round(ytd / months_elapsed, 2),
                  "projection": round(ytd + (ytd / months_elapsed) * (12 - months_elapsed), 2)}
bud_fy27 = sum(b[3] for b in budgets if b[1] == "FY2026-27" and b[2] in PRODUCT_CATS)
F["budget_fy27"] = round(bud_fy27, 2)
F["fy27_proj_util"] = round(F["fy2026_27"]["projection"] / bud_fy27, 3) if bud_fy27 else None

# requirement proposal for a few critical items
req = []
under_90 = [k for k, v in sorted(cover.items(), key=lambda x: x[1]) if v < 90]
for iid in under_90[:10]:
    recent = recent_item.get(iid, 0)
    daily = recent / 91.0
    target_cover = 60
    onhand = onhand_item.get(iid, 0)
    proposal = max(0, int(round(daily * target_cover - onhand)))
    req.append({"id": iid, "name": itm[iid][1], "cover": round(cover.get(iid, 0)),
                "onhand": round(onhand), "proposal": proposal,
                "unit": itm[iid][4], "cost": round(proposal * itm[iid][6], 2)})
F["requirement_proposal"] = req

# ---------------------------------------------------------------- commentary seeds
F["narrative"] = {
    "monsoon": "monsoon / diarrhoeal-disease season (Jul-Sep) drives ORS, IV fluid and antibiotic demand",
    "funding": "a funding interruption in Mar-2026 suppressed consumption for one month",
}

with open(OUT, "w") as f:
    json.dump(F, f, indent=2, default=str)

print(json.dumps({k: F[k] for k in [
    "consumption", "purchase", "ratio", "utilisation", "unattributed_pct",
    "peak_month", "trough_month", "yoy_growth", "top_item", "concentration",
    "price_items_measured", "avg_price_change", "cover", "stock", "fy2026_27"]}, indent=2, default=str)[:4000])
print("\nsaved", OUT)
